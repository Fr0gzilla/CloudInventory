import csv
import io
import os
from html.parser import HTMLParser

import pytest
from openpyxl import load_workbook

from app import db
from app.models import Account, Anomaly, Asset, ConsolidatedAsset, Run


def workbook(response):
    return load_workbook(io.BytesIO(response.data)).active


@pytest.fixture
def inventory(orm_app):
    run = Run(status="SUCCESS", started_at="2026-01-01T00:00:00Z")
    db.session.add(run)
    db.session.flush()
    for number in range(27):
        asset = Asset(run_id=run.id, vm_id=f"vm-{number:02d}", vm_name=f"VM {number:02d}",
                      node="special" if number == 0 else "common", type="qemu")
        db.session.add(asset)
        db.session.flush()
        db.session.add(ConsolidatedAsset(run_id=run.id, asset_id=asset.id,
                                         match_status="NO_MATCH" if number == 0 else "MATCHED_NAME"))
        if number == 0:
            db.session.add(Anomaly(run_id=run.id, asset_id=asset.id, code="NO_MATCH",
                                   description="missing", created_at="2026-01-01", severity="critique"))
    db.session.commit()
    client = orm_app.test_client()
    token = client.post("/api/login", json={"username": "admin", "password": os.environ["ADMIN_PASSWORD"]}).get_json()["access_token"]
    client.get("/login")
    with client.session_transaction() as session:
        csrf = session["csrf_token"]
    assert client.post("/login", data={"username": "admin", "password": os.environ["ADMIN_PASSWORD"],
                                      "csrf_token": csrf}).status_code == 302
    return client, {"Authorization": f"Bearer {token}"}, run


@pytest.mark.parametrize("prefix", ["/api/inventory", "/inventory"])
def test_excel_matches_filtered_sorted_csv_without_pagination(inventory, prefix):
    client, headers, _ = inventory
    headers = headers if prefix.startswith("/api/") else {}
    query = "?node=common&sort=vm_id&order=desc"
    csv_response = client.get(prefix + "/export" + query, headers=headers)
    excel_response = client.get(prefix + "/export.xlsx" + query, headers=headers)
    expected = list(csv.reader(io.StringIO(csv_response.get_data(as_text=True))))
    sheet = workbook(excel_response)
    actual = [["" if value is None else str(value) for value in row]
              for row in sheet.iter_rows(values_only=True)]
    assert (excel_response.status_code, len(actual), actual) == (200, 27, expected)


def test_excel_has_filter_frozen_headers_readable_widths_and_anomaly_row_color(inventory):
    client, headers, _ = inventory
    sheet = workbook(client.get("/api/inventory/export.xlsx?node=special", headers=headers))
    assert sheet.freeze_panes == "A2"
    assert sheet.auto_filter.ref == f"A1:Y{sheet.max_row}"
    assert all(sheet.column_dimensions[cell.column_letter].width > len(str(cell.value))
               for cell in sheet[1])
    assert sheet["A2"].value == "vm-00"
    assert sheet["A2"].fill.fgColor.rgb == "00F4CCCC" or sheet["A2"].fill.fgColor.rgb == "FFF4CCCC"
    assert sheet["Y2"].fill.fgColor.rgb == sheet["A2"].fill.fgColor.rgb


@pytest.mark.parametrize("prefix", ["=", "+", "-", "@"])
def test_excel_neutralizes_formula_prefixes_in_each_cell(inventory, prefix):
    client, headers, run = inventory
    asset = Asset(run_id=run.id, vm_id=f"{prefix}formula", vm_name=f"{prefix}name", type="qemu")
    db.session.add(asset)
    db.session.flush()
    db.session.add(ConsolidatedAsset(run_id=run.id, asset_id=asset.id, match_status="NO_MATCH"))
    db.session.commit()
    sheet = workbook(client.get("/api/inventory/export.xlsx?sort=vm_id", headers=headers))
    row = next(row for row in sheet.iter_rows() if row[0].value == f"'{prefix}formula")
    assert [(cell.value, cell.data_type) for cell in row[:2]] == [
        (f"'{prefix}formula", "s"), (f"'{prefix}name", "s")]


@pytest.mark.parametrize("prefix", ["/api/inventory", "/inventory"])
def test_excel_empty_inventory_contains_only_csv_headers(orm_app, prefix):
    client = orm_app.test_client()
    if prefix.startswith("/api/"):
        token = client.post("/api/login", json={"username": "admin", "password": os.environ["ADMIN_PASSWORD"]}).get_json()["access_token"]
        headers = {"Authorization": f"Bearer {token}"}
    else:
        client.get("/login")
        with client.session_transaction() as session:
            csrf = session["csrf_token"]
        client.post("/login", data={"username": "admin", "password": os.environ["ADMIN_PASSWORD"], "csrf_token": csrf})
        headers = {}
    sheet = workbook(client.get(prefix + "/export.xlsx", headers=headers))
    csv_rows = list(csv.reader(io.StringIO(client.get(prefix + "/export", headers=headers).get_data(as_text=True))))
    assert list(sheet.values) == [tuple(csv_rows[0])]


@pytest.mark.parametrize("prefix", ["/api/inventory", "/inventory"])
@pytest.mark.parametrize("query", ["?page=1", "?per_page=2", "?sort=invalid", "?match=WRONG"])
def test_excel_rejects_invalid_filters_and_pagination(inventory, prefix, query):
    client, headers, _ = inventory
    response = client.get(prefix + "/export.xlsx" + query,
                          headers=headers if prefix.startswith("/api/") else {})
    assert (response.status_code, response.get_json()) == (400, {"error": "Invalid inventory query"})


@pytest.mark.parametrize("headers", [{}, {"Authorization": "Bearer invalid"}])
def test_excel_api_requires_valid_bearer(orm_app, headers):
    response = orm_app.test_client().get("/api/inventory/export.xlsx", headers=headers)
    assert (response.status_code, response.get_json()) == (401, {"error": "Unauthorized"})


def test_excel_session_redirects_anonymous_user(orm_app):
    response = orm_app.test_client().get("/inventory/export.xlsx")
    assert response.status_code == 302 and "/login" in response.headers["Location"]


@pytest.mark.parametrize("role", ["lecteur", "opérateur"])
def test_excel_read_roles_can_download_but_disabled_account_cannot(inventory, role):
    client, _, _ = inventory
    client.get("/accounts")
    with client.session_transaction() as session:
        csrf = session["csrf_token"]
    assert client.post("/accounts", data={"username": role, "password": "test-password-123",
                                          "role": role, "csrf_token": csrf}).status_code == 302
    reader = client.application.test_client()
    reader.get("/login")
    with reader.session_transaction() as session:
        reader_csrf = session["csrf_token"]
    assert reader.post("/login", data={"username": role, "password": "test-password-123",
                                       "csrf_token": reader_csrf}).status_code == 302
    assert reader.get("/inventory/export.xlsx").status_code == 200
    account = Account.query.filter_by(username=role).one()
    assert client.post(f"/accounts/{account.id}", data={"action": "disable", "csrf_token": csrf}).status_code == 302
    assert reader.get("/inventory/export.xlsx").status_code != 200


class ExportLinks(HTMLParser):
    def __init__(self):
        super().__init__()
        self.links = {}

    def handle_starttag(self, tag, attrs):
        attrs = dict(attrs)
        if tag == "a" and attrs.get("id") in ("inventory-export", "inventory-export-excel"):
            self.links[attrs["id"]] = attrs.get("href")


def test_inventory_page_preserves_csv_and_exposes_excel_button(inventory):
    client, _, _ = inventory
    markup = ExportLinks()
    markup.feed(client.get("/inventory").get_data(as_text=True))
    assert markup.links == {"inventory-export": "/inventory/export",
                            "inventory-export-excel": "/inventory/export.xlsx"}


def test_inventory_script_passes_same_filter_parameters_to_both_exports():
    from pathlib import Path

    script = Path(__file__).resolve().parents[1].joinpath("app/static/inventory.js").read_text()
    assert "exportLink.href = `/inventory/export?${params}`;" in script
    assert "excelExportLink.href = `/inventory/export.xlsx?${params}`;" in script
